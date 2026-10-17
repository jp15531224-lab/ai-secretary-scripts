#!/usr/bin/env python3
"""食事DB（Googleスプレッドシート）の初期ファイル(.xlsx)を作る。Drive にアップロードすると Google 形式に変換される。

  python3 build_db_template.py <出力.xlsx> [--from 2026-10-15] [--to 2027-06-30]

シートの列は meal_core.py の定義と同じ。「日次集計」は人が見る用の数式シート（Make・Claude は使わない）。
"""
import argparse
from datetime import datetime, timedelta

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

import meal_core as mc

MAX_ROW = 3000  # 日次集計が参照する「食事」シートの行数（足りなくなったら数式の範囲を広げる）
HEAD_FILL = PatternFill("solid", fgColor="DDEBF7")


def header(ws, cols, widths=None):
    ws.append(cols)
    for i, _ in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i)
        c.font = Font(bold=True)
        c.fill = HEAD_FILL
        ws.column_dimensions[get_column_letter(i)].width = (widths or {}).get(i, 12)
    ws.freeze_panes = "A2"


def build(path, start, end, daily=True):
    wb = Workbook()
    first = True
    for name, cols in mc.SHEETS.items():
        ws = wb.active if first else wb.create_sheet(name)
        ws.title = name
        first = False
        header(ws, cols, {1: 22, 2: 20, 4: 17, 7: 22} if name == "食事" else None)
        if name == "設定":
            ws.column_dimensions["A"].width = 16
            ws.column_dimensions["C"].width = 70
            for row in mc.DEFAULT_SETTINGS:
                ws.append(row)

    if not daily:
        wb.save(path)
        return 0
    ws = wb.create_sheet("日次集計", 0)
    cols = ["集計日", "曜日", "品数", "kcal", "食塩g", "たんぱく質g", "脂質g", "炭水化物g", "アルコールg",
            "目標kcal", "差（kcal−目標）", "7日平均kcal", "7日平均食塩g"]
    header(ws, cols)
    ws.cell(row=1, column=14, value="午前6時〜翌5:59を1日として集計（推定値）。状態が「確定」の品目×食べた割合の合計。"
                                    "7日平均は記録のある日だけの平均。このシートは閲覧用です。")
    rng = lambda col: f"'食事'!${col}$2:${col}${MAX_ROW}"  # noqa: E731
    cond = f"({rng('E')}=$A{{r}})*({rng('C')}=\"確定\")"
    d = start
    r = 2
    while d <= end:
        key = d.strftime("%Y-%m-%d")
        ws.cell(row=r, column=1, value=key)
        ws.cell(row=r, column=2, value=mc.WEEKDAYS[d.weekday()])
        ws.cell(row=r, column=3, value=f"=SUMPRODUCT({cond.format(r=r)})")
        for i, col in enumerate("JMNOPQ", start=4):
            ws.cell(row=r, column=i, value=f"=ROUND(SUMPRODUCT({cond.format(r=r)},{rng(col)},{rng('I')}),1)")
        ws.cell(row=r, column=10, value="='設定'!$B$2")
        ws.cell(row=r, column=11, value=f'=IF(C{r}=0,"",D{r}-J{r})')
        lo = max(2, r - 6)
        ws.cell(row=r, column=12, value=f'=IFERROR(ROUND(AVERAGEIF(D{lo}:D{r},">0"),0),"")')
        ws.cell(row=r, column=13, value=f'=IFERROR(ROUND(AVERAGEIFS(E{lo}:E{r},D{lo}:D{r},">0"),1),"")')
        d += timedelta(days=1)
        r += 1
    wb.save(path)
    return r - 2


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("out")
    ap.add_argument("--from", dest="start", default="2026-10-15")
    ap.add_argument("--to", dest="end", default="2027-06-30")
    ap.add_argument("--no-daily", action="store_true", help="日次集計シートを作らない（後で Sheets API で追加する場合）")
    a = ap.parse_args()
    n = build(a.out, datetime.strptime(a.start, "%Y-%m-%d"), datetime.strptime(a.end, "%Y-%m-%d"), not a.no_daily)
    print(f"ok: {a.out}（日次集計 {n} 日分）")
